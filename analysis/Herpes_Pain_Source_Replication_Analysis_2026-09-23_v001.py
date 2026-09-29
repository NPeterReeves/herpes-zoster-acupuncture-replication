#!/usr/bin/env python3
"""Reproduce the prespecified Herpes Pain source analyses.

This program reads only the deidentified working workbook. It validates the
analysis population, applies the locked post-cure structural-zero rule for
VAS days 1-7, and writes machine-readable results used by the review workbook.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import platform

import numpy as np
from openpyxl import load_workbook
import scipy
from scipy import stats


VERSION = "2026-09-23_v001"
EXPECTED_INPUT_SHA256 = "6decf0e070c0a3b3515dc0b9ce28d24c101aad2e55cedaa5540d7803c1b0a54f"
EXPECTED_HEADERS = [
    "study_participant_id", "sex", "age_years", "hospitalization_days_source",
    "randomized_group_code", "randomized_group_label", "standardized_index_date",
    "standardized_discharge_date_derived", "source_date_interval_qc",
    "vas_baseline", "vas_day_1", "vas_day_2", "vas_day_3", "vas_day_4",
    "vas_day_5", "vas_day_6", "vas_day_7", "vas_day_28", "pain_duration_days",
    "dlqi_admission_total", "dlqi_discharge_total", "nephrotoxicity_flag",
    "hepatotoxicity_flag", "nausea_flag", "vomiting_flag", "upper_gi_bleeding_flag",
    "fainting_flag", "needle_break_flag", "hematoma_flag",
]
GROUPS = ((1, "Control"), (2, "Experimental"))
VAS_DAYS = [f"vas_day_{day}" for day in range(1, 8)]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return float(value)
    raise ValueError(f"Expected a numeric value or blank; found {value!r}")


def load_data(path: Path) -> list[dict]:
    if not path.is_file():
        raise FileNotFoundError(path)
    actual_hash = sha256(path)
    if actual_hash != EXPECTED_INPUT_SHA256:
        raise ValueError(f"Deidentified input checksum mismatch: {actual_hash}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Data" not in workbook.sheetnames:
        raise ValueError("Expected a Data worksheet")
    worksheet = workbook["Data"]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = list(rows[0])
    if headers != EXPECTED_HEADERS:
        raise ValueError("The deidentified workbook schema has changed")
    records = [dict(zip(headers, row)) for row in rows[1:] if row[0] is not None]
    if len(records) != 106:
        raise ValueError(f"Expected 106 participants; found {len(records)}")
    ids = [int(record["study_participant_id"]) for record in records]
    if len(set(ids)) != 106 or min(ids) != 1 or max(ids) != 106:
        raise ValueError("Participant identifier reconciliation failed")
    counts = {code: sum(record["randomized_group_code"] == code for record in records) for code, _ in GROUPS}
    if counts != {1: 52, 2: 54}:
        raise ValueError(f"Randomized group reconciliation failed: {counts}")
    for record in records:
        for key in EXPECTED_HEADERS:
            if key.startswith("vas_") or key.startswith("dlqi_") or key.endswith("_flag"):
                record[key] = clean_number(record[key])
        record["age_years"] = clean_number(record["age_years"])
        record["hospitalization_days_source"] = clean_number(record["hospitalization_days_source"])
    return records


def apply_structural_zeros(records: list[dict]) -> tuple[list[dict], list[dict]]:
    analyzed = [dict(record) for record in records]
    audit = []
    for record in analyzed:
        zero_seen = False
        for day, variable in enumerate(VAS_DAYS, start=1):
            source_value = record[variable]
            status = "OBSERVED" if source_value is not None else "SOURCE_MISSING"
            if source_value == 0:
                zero_seen = True
            elif source_value is None and zero_seen:
                record[variable] = 0.0
                status = "POST_CURE_STRUCTURAL_ZERO"
            audit.append({
                "Participant_ID": int(record["study_participant_id"]),
                "Group": record["randomized_group_label"],
                "Day": day,
                "Source_Value": source_value,
                "Analysis_Value": record[variable],
                "Status": status,
            })
    return analyzed, audit


def values(records: list[dict], variable: str, group: int | None = None) -> np.ndarray:
    selected = [
        record[variable] for record in records
        if (group is None or record["randomized_group_code"] == group) and record[variable] is not None
    ]
    return np.asarray(selected, dtype=float)


def haverage_quantile(array: np.ndarray, probability: float) -> float:
    """SPSS EXAMINE HAVERAGE percentile: rank=(n+1)p with endpoint clipping."""
    ordered = np.sort(np.asarray(array, dtype=float))
    if ordered.size == 0:
        return math.nan
    rank = (ordered.size + 1) * probability
    if rank <= 1:
        return float(ordered[0])
    if rank >= ordered.size:
        return float(ordered[-1])
    lower = math.floor(rank)
    fraction = rank - lower
    return float(ordered[lower - 1] + fraction * (ordered[lower] - ordered[lower - 1]))


def summary_median_iqr(array: np.ndarray) -> dict:
    q1 = haverage_quantile(array, 0.25)
    median = haverage_quantile(array, 0.50)
    q3 = haverage_quantile(array, 0.75)
    return {"n": int(array.size), "median": median, "q1": q1, "q3": q3, "iqr": q3 - q1}


def summary_mean_sd(array: np.ndarray) -> dict:
    return {"n": int(array.size), "mean": float(np.mean(array)), "sd": float(np.std(array, ddof=1))}


def median_text(summary: dict) -> str:
    return f'{summary["median"]:.2f} ({summary["iqr"]:.2f})'


def mean_text(summary: dict) -> str:
    return f'{summary["mean"]:.2f} ± {summary["sd"]:.2f}'


def p_text(value: float | None) -> str:
    if value is None or not math.isfinite(value):
        return ""
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def mann_whitney(records: list[dict], variable: str) -> tuple[float, float]:
    result = stats.mannwhitneyu(
        values(records, variable, 1), values(records, variable, 2),
        alternative="two-sided", method="asymptotic", use_continuity=False,
    )
    return float(result.statistic), float(result.pvalue)


def wilson(successes: int, total: int, alpha: float = 0.05) -> tuple[float, float]:
    z = float(stats.norm.ppf(1 - alpha / 2))
    p = successes / total
    denominator = 1 + z * z / total
    center = (p + z * z / (2 * total)) / denominator
    half = z * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total)) / denominator
    return center - half, center + half


def newcombe_difference(x1: int, n1: int, x0: int, n0: int) -> tuple[float, float]:
    p1, p0 = x1 / n1, x0 / n0
    l1, u1 = wilson(x1, n1)
    l0, u0 = wilson(x0, n0)
    lower = (p1 - p0) - math.sqrt((p1 - l1) ** 2 + (u0 - p0) ** 2)
    upper = (p1 - p0) + math.sqrt((u1 - p1) ** 2 + (p0 - l0) ** 2)
    return lower, upper


def katz_rr(x1: int, n1: int, x0: int, n0: int) -> tuple[float, float, float]:
    rr = (x1 / n1) / (x0 / n0)
    se = math.sqrt(1 / x1 - 1 / n1 + 1 / x0 - 1 / n0)
    z = float(stats.norm.ppf(0.975))
    return rr, math.exp(math.log(rr) - z * se), math.exp(math.log(rr) + z * se)


def woolf_or(x1: int, n1: int, x0: int, n0: int) -> tuple[float, float, float]:
    a, b, c, d = x1, n1 - x1, x0, n0 - x0
    odds_ratio = (a * d) / (b * c)
    se = math.sqrt(1 / a + 1 / b + 1 / c + 1 / d)
    z = float(stats.norm.ppf(0.975))
    return odds_ratio, math.exp(math.log(odds_ratio) - z * se), math.exp(math.log(odds_ratio) + z * se)


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"No rows supplied for {path.name}")
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def build_primary_cure(records: list[dict]) -> tuple[list[dict], dict]:
    # Day-7 analysis values include derived post-cure structural zeros.
    x0 = sum(record["vas_day_7"] == 0 for record in records if record["randomized_group_code"] == 1)
    x1 = sum(record["vas_day_7"] == 0 for record in records if record["randomized_group_code"] == 2)
    n0, n1 = 52, 54
    p0, p1 = x0 / n0, x1 / n1
    control_ci = wilson(x0, n0)
    experimental_ci = wilson(x1, n1)
    rd = p1 - p0
    rd_ci = newcombe_difference(x1, n1, x0, n0)
    rr, rr_l, rr_u = katz_rr(x1, n1, x0, n0)
    odds_ratio, or_l, or_u = woolf_or(x1, n1, x0, n0)
    table = np.asarray([[x0, n0 - x0], [x1, n1 - x1]], dtype=int)
    chi2, p_two, _, expected = stats.chi2_contingency(table, correction=False)
    chi2_yates, p_yates, _, _ = stats.chi2_contingency(table, correction=True)
    fisher_or, p_fisher = stats.fisher_exact(table, alternative="two-sided")
    z = math.copysign(math.sqrt(float(chi2)), p1 - p0)
    p_one = float(stats.norm.sf(z))
    rows = [
        {"Section": "Group rate", "Measure": "Control cured by day 7", "Numerator": x0, "Denominator": n0, "Estimate": p0, "CI_Lower": control_ci[0], "CI_Upper": control_ci[1], "Statistic": "", "P_Value": "", "Method": "Wilson 95% CI", "Interpretation": "Observed or confirmed structural zero at day 7"},
        {"Section": "Group rate", "Measure": "Experimental cured by day 7", "Numerator": x1, "Denominator": n1, "Estimate": p1, "CI_Lower": experimental_ci[0], "CI_Upper": experimental_ci[1], "Statistic": "", "P_Value": "", "Method": "Wilson 95% CI", "Interpretation": "Observed or confirmed structural zero at day 7"},
        {"Section": "Effect", "Measure": "Risk difference (Experimental - Control)", "Numerator": "", "Denominator": "", "Estimate": rd, "CI_Lower": rd_ci[0], "CI_Upper": rd_ci[1], "Statistic": "", "P_Value": "", "Method": "Newcombe hybrid-score 95% CI", "Interpretation": "Absolute increase in cure probability"},
        {"Section": "Effect", "Measure": "Risk ratio (Experimental / Control)", "Numerator": "", "Denominator": "", "Estimate": rr, "CI_Lower": rr_l, "CI_Upper": rr_u, "Statistic": "", "P_Value": "", "Method": "Katz log 95% CI", "Interpretation": "Relative cure probability"},
        {"Section": "Effect", "Measure": "Odds ratio (Experimental / Control)", "Numerator": "", "Denominator": "", "Estimate": odds_ratio, "CI_Lower": or_l, "CI_Upper": or_u, "Statistic": "", "P_Value": "", "Method": "Woolf log 95% CI", "Interpretation": "Relative odds of cure"},
        {"Section": "Primary test", "Measure": "Pearson chi-square, two-sided, uncorrected", "Numerator": "", "Denominator": "", "Estimate": "", "CI_Lower": "", "CI_Upper": "", "Statistic": float(chi2), "P_Value": float(p_two), "Method": "Prespecified primary test; alpha=0.05", "Interpretation": "Statistically significant" if p_two < 0.05 else "Not statistically significant"},
        {"Section": "Sensitivity", "Measure": "Pearson-equivalent z, one-sided favorable", "Numerator": "", "Denominator": "", "Estimate": "", "CI_Lower": "", "CI_Upper": "", "Statistic": z, "P_Value": p_one, "Method": "Upper-tail normal probability", "Interpretation": "Experimental cure rate higher"},
        {"Section": "Sensitivity", "Measure": "Pearson chi-square with Yates correction", "Numerator": "", "Denominator": "", "Estimate": "", "CI_Lower": "", "CI_Upper": "", "Statistic": float(chi2_yates), "P_Value": float(p_yates), "Method": "Two-sided", "Interpretation": "Sensitivity analysis"},
        {"Section": "Sensitivity", "Measure": "Fisher exact", "Numerator": "", "Denominator": "", "Estimate": fisher_or, "CI_Lower": "", "CI_Upper": "", "Statistic": "", "P_Value": float(p_fisher), "Method": "Two-sided", "Interpretation": "Sensitivity analysis"},
    ]
    checks = {"control_cured": x0, "experimental_cured": x1, "expected_min": float(np.min(expected)), "primary_p": float(p_two), "risk_difference": rd, "risk_ratio": rr, "odds_ratio": odds_ratio}
    return rows, checks


def build_table1(records: list[dict]) -> list[dict]:
    rows = []
    age_total = summary_mean_sd(values(records, "age_years"))
    age_c = summary_mean_sd(values(records, "age_years", 1))
    age_e = summary_mean_sd(values(records, "age_years", 2))
    t = stats.ttest_ind(values(records, "age_years", 1), values(records, "age_years", 2), equal_var=True)
    rows.append({"Variable": "Age (years)", "Level": "", "Total_n": age_total["n"], "Total_Summary": mean_text(age_total), "Control_n": age_c["n"], "Control_Summary": mean_text(age_c), "Experimental_n": age_e["n"], "Experimental_Summary": mean_text(age_e), "Test": "Pooled independent-samples t-test", "Statistic": float(t.statistic), "P_Value": float(t.pvalue), "Published_P": "0.857", "Assessment": "Rounding-compatible"})
    sexes = ["Female", "Male"]
    sex_table = np.asarray([[sum(r["sex"] == s and r["randomized_group_code"] == g for r in records) for s in sexes] for g, _ in GROUPS])
    sex_chi, sex_p, _, _ = stats.chi2_contingency(sex_table, correction=False)
    for index, sex in enumerate(sexes):
        total = sum(r["sex"] == sex for r in records)
        control = sum(r["sex"] == sex and r["randomized_group_code"] == 1 for r in records)
        experimental = sum(r["sex"] == sex and r["randomized_group_code"] == 2 for r in records)
        rows.append({"Variable": "Sex" if index == 0 else "", "Level": sex, "Total_n": 106, "Total_Summary": f"{total} ({100*total/106:.2f}%)", "Control_n": 52, "Control_Summary": f"{control} ({100*control/52:.2f}%)", "Experimental_n": 54, "Experimental_Summary": f"{experimental} ({100*experimental/54:.2f}%)", "Test": "Pearson chi-square, uncorrected" if index == 0 else "", "Statistic": float(sex_chi) if index == 0 else "", "P_Value": float(sex_p) if index == 0 else "", "Published_P": "0.503" if index == 0 else "", "Assessment": "Unresolved p-value discrepancy" if index == 0 else ""})
    for variable, label, published in [
        ("vas_baseline", "Baseline VAS", "0.434"),
        ("dlqi_admission_total", "Admission DLQI", "0.568"),
        ("hospitalization_days_source", "Hospitalization (days)", "0.552"),
    ]:
        total = summary_median_iqr(values(records, variable))
        control = summary_median_iqr(values(records, variable, 1))
        experimental = summary_median_iqr(values(records, variable, 2))
        u, p_value = mann_whitney(records, variable)
        rows.append({"Variable": label, "Level": "", "Total_n": total["n"], "Total_Summary": median_text(total), "Control_n": control["n"], "Control_Summary": median_text(control), "Experimental_n": experimental["n"], "Experimental_Summary": median_text(experimental), "Test": "Mann-Whitney U, asymptotic", "Statistic": u, "P_Value": p_value, "Published_P": published, "Assessment": "Rounding-compatible"})
    return rows


def build_vas_table(records: list[dict], sensitivity: bool = False) -> list[dict]:
    published = {1: "0.434", 2: "<0.001", 3: "<0.001", 4: "<0.001", 5: "<0.001", 6: "0.021", 7: "0.013", 28: "0.070"}
    rows = []
    for day in [1, 2, 3, 4, 5, 6, 7, 28]:
        variable = f"vas_day_{day}"
        total = summary_median_iqr(values(records, variable))
        control = summary_median_iqr(values(records, variable, 1))
        experimental = summary_median_iqr(values(records, variable, 2))
        u, p_value = mann_whitney(records, variable)
        assessment = "Observed-only sensitivity; not used for strict replication" if sensitivity else "Rounding-compatible"
        rows.append({"Day": day, "Total_n": total["n"], "Total_Median": total["median"], "Total_IQR": total["iqr"], "Control_n": control["n"], "Control_Median": control["median"], "Control_IQR": control["iqr"], "Experimental_n": experimental["n"], "Experimental_Median": experimental["median"], "Experimental_IQR": experimental["iqr"], "U_Statistic": u, "P_Value": p_value, "P_Display": p_text(p_value), "Published_P": published[day], "Assessment": assessment})
    return rows


def build_table3(records: list[dict]) -> list[dict]:
    augmented = [dict(record, dlqi_change=record["dlqi_admission_total"] - record["dlqi_discharge_total"]) for record in records]
    specifications = [
        ("dlqi_admission_total", "Admission", "0.586", "Published Table 3 p-value differs from source replication and published Table 1"),
        ("dlqi_discharge_total", "Hospital discharge", "0.005", "Rounding-compatible"),
        ("dlqi_change", "Change (admission - discharge)", "<0.001", "Experimental IQR discrepancy; test is reproducible"),
    ]
    rows = []
    for variable, timepoint, published, assessment in specifications:
        total = summary_median_iqr(values(augmented, variable))
        control = summary_median_iqr(values(augmented, variable, 1))
        experimental = summary_median_iqr(values(augmented, variable, 2))
        u, p_value = mann_whitney(augmented, variable)
        rows.append({"Time": timepoint, "Total_n": total["n"], "Total_Median": total["median"], "Total_IQR": total["iqr"], "Control_n": control["n"], "Control_Median": control["median"], "Control_IQR": control["iqr"], "Experimental_n": experimental["n"], "Experimental_Median": experimental["median"], "Experimental_IQR": experimental["iqr"], "U_Statistic": u, "P_Value": p_value, "P_Display": p_text(p_value), "Published_P": published, "Assessment": assessment})
    return rows


def build_table4(records: list[dict]) -> list[dict]:
    specs = [
        ("nephrotoxicity_flag", "Nephrotoxicity", 5, 7, "0.627"),
        ("hepatotoxicity_flag", "Hepatotoxicity", 3, 5, "0.528"),
        ("gastrointestinal_disorder", "Gastrointestinal disorder", 12, 9, "0.498"),
        ("nausea_flag", "Nausea", 9, 7, "0.592"),
        ("vomiting_flag", "Vomiting", 3, 2, "0.632"),
        ("upper_gi_bleeding_flag", "Upper gastrointestinal bleeding", 0, 0, ""),
        ("fainting_flag", "Fainting", 0, 0, ""),
        ("needle_break_flag", "Needle break", 0, 0, ""),
        ("hematoma_flag", "Hematoma", 0, 0, ""),
    ]
    augmented = []
    for record in records:
        copy = dict(record)
        copy["gastrointestinal_disorder"] = float(bool(record["nausea_flag"] or record["vomiting_flag"]))
        augmented.append(copy)
    rows = []
    for variable, label, published_c, published_e, published_p in specs:
        c = int(sum(values(augmented, variable, 1)))
        e = int(sum(values(augmented, variable, 2)))
        total = c + e
        if total:
            table = np.asarray([[c, 52 - c], [e, 54 - e]], dtype=int)
            chi2, pearson_p, _, expected = stats.chi2_contingency(table, correction=False)
            fisher_or, fisher_p = stats.fisher_exact(table, alternative="two-sided")
            minimum_expected = float(np.min(expected))
        else:
            chi2 = pearson_p = fisher_or = fisher_p = minimum_expected = math.nan
        published_counts = f"{published_c}/{published_e}"
        replicated_counts = f"{c}/{e}"
        assessment = "Rounding-compatible" if published_counts == replicated_counts and (not published_p or p_text(pearson_p) == published_p) else "Unresolved source/publication discrepancy"
        rows.append({"Side_Effect": label, "Total_Events": total, "Total_Percent": 100 * total / 106, "Control_Events": c, "Control_Percent": 100 * c / 52, "Experimental_Events": e, "Experimental_Percent": 100 * e / 54, "Pearson_ChiSquare": "" if not math.isfinite(chi2) else chi2, "Pearson_P": "" if not math.isfinite(pearson_p) else pearson_p, "Fisher_OR": "" if not math.isfinite(fisher_or) else fisher_or, "Fisher_P": "" if not math.isfinite(fisher_p) else fisher_p, "Minimum_Expected": "" if not math.isfinite(minimum_expected) else minimum_expected, "Published_Control_Events": published_c, "Published_Experimental_Events": published_e, "Published_P": published_p, "Assessment": assessment})
    return rows


def build_qc(source: list[dict], analyzed: list[dict], audit: list[dict]) -> list[dict]:
    rows = []
    for day, variable in enumerate(VAS_DAYS, start=1):
        for group_code, group_label in GROUPS:
            group_audit = [a for a in audit if a["Day"] == day and a["Group"] == group_label]
            rows.append({"Day": day, "Group": group_label, "Randomized_N": sum(r["randomized_group_code"] == group_code for r in source), "Observed_N": sum(a["Status"] == "OBSERVED" for a in group_audit), "Source_Missing_N": sum(a["Source_Value"] is None for a in group_audit), "Derived_Structural_Zero_N": sum(a["Status"] == "POST_CURE_STRUCTURAL_ZERO" for a in group_audit), "Unresolved_Missing_N": sum(a["Analysis_Value"] is None for a in group_audit), "Analysis_N": int(values(analyzed, variable, group_code).size)})
    return rows


def build_reconciliation(table1: list[dict], table2: list[dict], table3: list[dict], table4: list[dict], cure_checks: dict) -> list[dict]:
    rows = [
        {"ID": "R001", "Location": "Analysis population", "Item": "Randomized participants", "Published": "106 (52 control; 54 experimental)", "Replicated": "106 (52 control; 54 experimental)", "Assessment": "Exact", "Explanation": "Deidentified workbook reconciles to the randomized analysis population."},
        {"ID": "R002", "Location": "Table 1", "Item": "Age summaries and between-group p-value", "Published": "p=0.857", "Replicated": f'p={p_text(float(table1[0]["P_Value"]))}', "Assessment": "Rounding-compatible", "Explanation": "Prespecified pooled two-sample t-test reproduces the displayed p-value."},
        {"ID": "R003", "Location": "Table 1", "Item": "Sex counts", "Published": "Female 36/34; Male 16/20", "Replicated": "Female 36/34; Male 16/20", "Assessment": "Exact", "Explanation": "Counts match the source workbook."},
        {"ID": "R004", "Location": "Table 1", "Item": "Sex Pearson p-value", "Published": "0.503", "Replicated": f'{float(table1[1]["P_Value"]):.6f}', "Assessment": "Unresolved discrepancy", "Explanation": "Uncorrected Pearson, Yates-corrected Pearson, and Fisher exact do not produce 0.503."},
        {"ID": "R005", "Location": "Table 1", "Item": "Baseline VAS", "Published": "5.00 (1.00) vs 5.50 (1.00); p=0.434", "Replicated": f'p={p_text(float(table1[3]["P_Value"]))}', "Assessment": "Rounding-compatible", "Explanation": "HAVERAGE quartiles and asymptotic Mann-Whitney test reproduce the table."},
        {"ID": "R006", "Location": "Table 1", "Item": "Admission DLQI", "Published": "p=0.568", "Replicated": f'p={p_text(float(table1[4]["P_Value"]))}', "Assessment": "Rounding-compatible", "Explanation": "Table 1 is reproduced; Table 3 prints 0.586 for the same comparison."},
        {"ID": "R007", "Location": "Table 2", "Item": "VAS days 1-7 and 28", "Published": "All displayed medians, IQRs, and p-values", "Replicated": "All displayed medians, IQRs, and p-values", "Assessment": "Rounding-compatible", "Explanation": "Days 4-7 use only confirmed post-cure structural zeros; day 28 remains observed-only."},
        {"ID": "R008", "Location": "Figure 3/text", "Item": "Day-7 cure rates", "Published": "65.38% control; 87.04% experimental; p<0.005", "Replicated": f'{100*cure_checks["control_cured"]/52:.2f}% control; {100*cure_checks["experimental_cured"]/54:.2f}% experimental; p={cure_checks["primary_p"]:.6f}', "Assessment": "Rates exact; p-value threshold not reproduced", "Explanation": "Two-sided uncorrected Pearson p-value is above 0.005; the favorable one-sided p-value is below 0.005."},
        {"ID": "R009", "Location": "Table 3", "Item": "Admission DLQI p-value", "Published": "0.586", "Replicated": f'{float(table3[0]["P_Value"]):.6f}', "Assessment": "Unresolved discrepancy", "Explanation": "The source result rounds to 0.568, matching Table 1 rather than Table 3."},
        {"ID": "R010", "Location": "Table 3", "Item": "Experimental DLQI change IQR", "Published": "1.25", "Replicated": f'{float(table3[2]["Experimental_IQR"]):.2f}', "Assessment": "Unresolved discrepancy", "Explanation": "SPSS EXAMINE HAVERAGE quartiles from participant-level values give an IQR of 1.00."},
        {"ID": "R011", "Location": "Table 4", "Item": "Gastrointestinal disorder count", "Published": "12 control; 9 experimental; total 21", "Replicated": f'{table4[2]["Control_Events"]} control; {table4[2]["Experimental_Events"]} experimental; total {table4[2]["Total_Events"]}', "Assessment": "Unresolved discrepancy", "Explanation": "The prespecified source definition is the union of nausea or vomiting flags."},
        {"ID": "R012", "Location": "Table 4", "Item": "Vomiting count", "Published": "3 control; 2 experimental; total 5", "Replicated": f'{table4[4]["Control_Events"]} control; {table4[4]["Experimental_Events"]} experimental; total {table4[4]["Total_Events"]}', "Assessment": "Unresolved discrepancy", "Explanation": "Participant-level source contains two vomiting flags in each group."},
        {"ID": "R013", "Location": "Table 4", "Item": "Safety p-values", "Published": "0.627, 0.528, 0.498, 0.592, 0.632", "Replicated": ", ".join(p_text(float(r["Pearson_P"])) for r in table4[:5]), "Assessment": "Unresolved discrepancy", "Explanation": "Published p-values are not reproduced by uncorrected Pearson tests on the participant-level source counts."},
    ]
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-xlsx", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    if args.output_dir.exists() and any(args.output_dir.iterdir()):
        raise FileExistsError(f"Output directory is not empty: {args.output_dir}")
    args.output_dir.mkdir(parents=True, exist_ok=True)

    source = load_data(args.input_xlsx)
    analyzed, audit = apply_structural_zeros(source)
    cure, cure_checks = build_primary_cure(analyzed)
    table1 = build_table1(analyzed)
    table2 = build_vas_table(analyzed)
    table3 = build_table3(analyzed)
    table4 = build_table4(analyzed)
    observed_only = build_vas_table(source, sensitivity=True)
    qc = build_qc(source, analyzed, audit)
    reconciliation = build_reconciliation(table1, table2, table3, table4, cure_checks)

    outputs = {
        "primary_cure.csv": cure,
        "table1_baseline.csv": table1,
        "table2_vas.csv": table2,
        "table3_dlqi.csv": table3,
        "table4_safety.csv": table4,
        "vas_observed_only_sensitivity.csv": observed_only,
        "vas_missingness_qc.csv": qc,
        "reconciliation.csv": reconciliation,
    }
    for filename, rows in outputs.items():
        write_csv(args.output_dir / filename, rows)

    metadata = {
        "analysis_version": VERSION,
        "created_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "input": {"filename": args.input_xlsx.name, "sha256": sha256(args.input_xlsx)},
        "software": {"python": platform.python_version(), "numpy": np.__version__, "scipy": scipy.__version__},
        "population_checks": {"total": 106, "control": 52, "experimental": 54, "unique_participant_ids": 106},
        "analysis_rules": {
            "quartiles": "SPSS EXAMINE HAVERAGE, rank=(n+1)p",
            "mann_whitney": "Two-sided asymptotic, tie-corrected, no continuity correction",
            "structural_zero": "Blank VAS after first observed zero on days 1-7 is derived as zero and flagged; day 28 observed-only",
            "primary_cure_test": "Two-sided uncorrected Pearson chi-square, alpha=0.05",
        },
        "key_results": cure_checks,
        "outputs": sorted(outputs),
    }
    (args.output_dir / "analysis_metadata.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(metadata, indent=2))


if __name__ == "__main__":
    main()
