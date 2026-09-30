#!/usr/bin/env python3
"""Exploratory aggregate first-week VAS burden analysis of the QiBo Herpes Pain RCT.

Reads the approved deidentified working workbook. Emits aggregate JSON only;
no participant-level rows are written. Requires openpyxl, numpy, scipy.
"""
import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy import stats

EXPECTED_SHA256 = '6decf0e070c0a3b3515dc0b9ce28d24c101aad2e55cedaa5540d7803c1b0a54f'
SEED = 20260930
BOOTSTRAP_REPLICATES = 10000


def read_approved_workbook(path):
    actual = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual != EXPECTED_SHA256:
        raise ValueError('Input workbook SHA-256 does not match the approved deidentified copy')
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = wb['Data'].iter_rows(values_only=True)
    headers = next(rows)
    required = ['randomized_group_code', 'vas_baseline'] + [f'vas_day_{d}' for d in range(1, 8)]
    if not all(key in headers for key in required):
        raise ValueError('Required analysis columns are absent')
    ix = {name: headers.index(name) for name in required}
    records = [row for row in rows if row[0] is not None]
    group = np.asarray([int(row[ix['randomized_group_code']]) for row in records])
    baseline = np.asarray([float(row[ix['vas_baseline']]) for row in records])
    scores = [[row[ix[f'vas_day_{d}']] for d in range(1, 8)] for row in records]
    if len(records) != 106 or sorted([np.sum(group == 1), np.sum(group == 2)]) != [52, 54]:
        raise ValueError('The randomized population or arm sizes do not reconcile')
    if set(group) != {1, 2}:
        raise ValueError('Unexpected arm code')
    derived = 0
    pre_cure_missing = 0
    for scores_i in scores:
        prior_observed_zero = False
        for j, value in enumerate(scores_i):
            if value is None:
                if prior_observed_zero:
                    scores_i[j] = 0.0
                    derived += 1
                else:
                    pre_cure_missing += 1
            elif float(value) == 0:
                prior_observed_zero = True
    scores = np.asarray(scores, dtype=float)
    if pre_cure_missing or np.isnan(scores).any():
        raise ValueError('Unresolved pre-cure missing VAS: an additional rule is required')
    if not np.array_equal(baseline, scores[:, 0]):
        raise ValueError('Baseline VAS is not identical to day-1 VAS as in the approved copy')
    if derived != 50:
        raise ValueError(f'Unexpected structural-zero count: {derived}')
    return group == 2, baseline, scores, derived


def welch_summary(y, experimental):
    control = y[~experimental]
    waa = y[experimental]
    n0, n1 = len(control), len(waa)
    var0, var1 = control.var(ddof=1), waa.var(ddof=1)
    delta = waa.mean() - control.mean()
    se = np.sqrt(var0/n0 + var1/n1)
    df = (var0/n0 + var1/n1)**2 / ((var0/n0)**2/(n0-1) + (var1/n1)**2/(n1-1))
    ci = stats.t.interval(.95, df, loc=delta, scale=se)
    sd_pooled = np.sqrt(((n0-1)*var0 + (n1-1)*var1)/(n0+n1-2))
    hedges_g = delta/sd_pooled * (1 - 3/(4*(n0+n1-2)-1))
    return dict(n_control=n0, n_waa=n1, control_mean=float(control.mean()), waa_mean=float(waa.mean()),
                control_sd=float(control.std(ddof=1)), waa_sd=float(waa.std(ddof=1)),
                waa_minus_control=float(delta), ci95=[float(ci[0]),float(ci[1])],
                welch_p=float(stats.ttest_ind(waa,control,equal_var=False).pvalue),
                hedges_g=float(hedges_g))


def adjusted_hc3(y, experimental, baseline):
    n=len(y)
    X=np.column_stack((np.ones(n), experimental.astype(float), baseline-baseline.mean()))
    inv=np.linalg.inv(X.T@X)
    beta=inv@X.T@y
    residual=y-X@beta
    leverage=np.einsum('ij,jk,ik->i',X,inv,X)
    meat=(X.T*(residual/(1-leverage))**2)@X
    se=np.sqrt(np.diag(inv@meat@inv))
    df=n-X.shape[1]
    ci=beta[1]+np.asarray([-1,1])*stats.t.ppf(.975,df)*se[1]
    p=2*stats.t.sf(abs(beta[1]/se[1]),df)
    return dict(waa_minus_control=float(beta[1]),ci95_hc3=[float(ci[0]),float(ci[1])],
                hc3_p=float(p), degrees_of_freedom=df)


def bootstrap_adjusted_auc(y, experimental, baseline):
    rng=np.random.default_rng(SEED)
    a=np.flatnonzero(~experimental)
    b=np.flatnonzero(experimental)
    estimates=[]
    for _ in range(BOOTSTRAP_REPLICATES):
        sel=np.r_[rng.choice(a,len(a),replace=True),rng.choice(b,len(b),replace=True)]
        X=np.column_stack((np.ones(len(sel)),experimental[sel].astype(float),baseline[sel]-baseline[sel].mean()))
        estimates.append(np.linalg.lstsq(X,y[sel],rcond=None)[0][1])
    q=np.quantile(estimates,[.025,.975])
    return [float(q[0]),float(q[1])]


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,required=True,help='Approved deidentified working XLSX')
    parser.add_argument('--output',type=Path,required=True,help='Aggregate-only JSON output path')
    args=parser.parse_args()
    experimental,baseline,scores,derived=read_approved_workbook(args.input)
    auc=np.trapezoid(scores,dx=1,axis=1)
    mean_daily=scores.mean(axis=1)
    mean_postbaseline=scores[:,1:].mean(axis=1)
    result={
      'analysis_status':'EXPLORATORY_POST_HOC',
      'input_sha256':EXPECTED_SHA256,
      'sample':{'n':106,'control':52,'waa_plus_medication':54,'derived_post_cure_zeros':derived,'unresolved_week_one_vas_missing':0},
      'definition':{'auc':'Trapezoidal area from day 1 to day 7 inclusive, six unit-day intervals; daily VAS reports worst pain in prior 24 hours. Units: VAS score-days.',
                    'baseline':'vas_baseline, numerically identical to vas_day_1 in all 106 records',
                    'contrast':'WAA plus medication minus medication alone; negative values favor WAA',
                    'primary_trial_endpoint_unchanged':'Day-7 VAS=0 cure, two-sided Pearson chi-square'},
      'auc_day1_to_day7':{'unadjusted_welch':welch_summary(auc,experimental),
                          'baseline_adjusted_ols_hc3':adjusted_hc3(auc,experimental,baseline),
                          'baseline_adjusted_stratified_bootstrap_ci95':bootstrap_adjusted_auc(auc,experimental,baseline)},
      'sensitivity_mean_of_seven_daily_vas':{'unadjusted_welch':welch_summary(mean_daily,experimental),
                                               'baseline_adjusted_ols_hc3':adjusted_hc3(mean_daily,experimental,baseline)},
      'sensitivity_mean_days2_to7':{'unadjusted_welch':welch_summary(mean_postbaseline,experimental)},
      'limits':['Daily VAS recorded worst pain in the prior 24 hours; the AUC is a trajectory summary, not observed time-integrated pain.',
                'Daily values after an observed zero were derived as zero through day 7 under the locked rule.',
                'The comparison includes WAA-specific and contextual effects because the trial has no sham or attention-matched arm.',
                'The analysis is post hoc; no multiple-analysis familywise correction is claimed.',
                'Actual medication consumption and exact symptom-onset dates are unavailable.']
    }
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'status':'PASS','output':str(args.output),'n':106,'auc_mean_control':result['auc_day1_to_day7']['unadjusted_welch']['control_mean'],'auc_mean_waa':result['auc_day1_to_day7']['unadjusted_welch']['waa_mean'],'adjusted_difference':result['auc_day1_to_day7']['baseline_adjusted_ols_hc3']['waa_minus_control']},indent=2))

if __name__=='__main__':main()
