#!/usr/bin/env python3
"""Create the approved deidentified Herpes Pain working workbook.

The script reads the restricted source workbook but never modifies it. It builds
a new workbook from an explicit allowlist so removed identifiers cannot remain
in hidden worksheets, deleted cells, shared strings, or workbook metadata.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timedelta
import hashlib
from pathlib import Path
import re
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as exc:  # pragma: no cover - environment preflight
    raise SystemExit("openpyxl is required. Install it in the project Python environment.") from exc


EXPECTED_SOURCE_SHA256 = "5e2f51bf004344bd3215be266c7115d8104cc26431ea89d5d9a047c0fc79edb3"
STANDARDIZED_INDEX_DATE = date(2000, 1, 1)
VERSION = "2026-09-23_v001"

HEADERS = [
    "study_participant_id", "sex", "age_years", "hospitalization_days_source",
    "randomized_group_code", "randomized_group_label", "standardized_index_date",
    "standardized_discharge_date_derived", "source_date_interval_qc",
    "vas_baseline", "vas_day_1", "vas_day_2", "vas_day_3", "vas_day_4",
    "vas_day_5", "vas_day_6", "vas_day_7", "vas_day_28", "pain_duration_days",
    "dlqi_admission_total", "dlqi_discharge_total", "nephrotoxicity_flag",
    "hepatotoxicity_flag", "nausea_flag", "vomiting_flag", "upper_gi_bleeding_flag",
    "fainting_flag", "needle_break_flag", "hematoma_flag",
]

DICTIONARY = [
    ["study_participant_id", "integer", "Number", "Sequential study participant identifier", "1-106; unique", "Copied from source sequential number", "Use to generate OMOP person_id"],
    ["sex", "text", "Gender", "Participant sex as supplied", "Female; Male", "Copied without recoding", "Map to OMOP gender concept during ETL"],
    ["age_years", "integer", "Age", "Age at study index in years", "25-80", "Copied without change", "Source age; do not infer an actual birth date"],
    ["hospitalization_days_source", "integer", "Hospitalization days", "Reported hospitalization duration", "4-22 days", "Copied without change", "Preserved as the source duration variable"],
    ["randomized_group_code", "integer", "Groups", "Randomized treatment group code", "1=Control; 2=Experimental", "Copied without change", "Randomized assignment, not confirmed treatment exposure"],
    ["randomized_group_label", "text", "Derived from Groups", "Readable randomized group label", "Control; Experimental", "1 mapped to Control; 2 mapped to Experimental", "Derived field"],
    ["standardized_index_date", "date", "Derived", "Synthetic study index date", "2000-01-01 for every participant", "Replaces the original admission date", "Contains no original calendar date"],
    ["standardized_discharge_date_derived", "date", "Derived", "Synthetic visit end date", "2000-01-04 through 2000-01-22", "Index date + hospitalization_days_source - 1", "Uses the predominant inclusive-day convention"],
    ["source_date_interval_qc", "text", "Derived from original dates and hospitalization days", "Reconciliation of the source date interval", "MATCH_INCLUSIVE; MATCH_ELAPSED; DISCREPANT; UNPARSABLE", "Contains the reconciliation category only; original dates are excluded", "Quality-control field; do not use as an outcome"],
    ["vas_baseline", "integer", "Baseline pain score", "Baseline visual analogue scale pain score", "0-10; observed source range 3-8", "Copied without change", "Baseline and day 1 are identical for all records"],
    ["vas_day_1", "integer", "Day 1 pain score", "Day 1 visual analogue scale pain score", "0-10; observed source range 3-8", "Copied without change", "Retained separately from baseline"],
    ["vas_day_2", "integer", "Day 2 pain score", "Day 2 visual analogue scale pain score", "0-10", "Copied without change", "Observed source value"],
    ["vas_day_3", "integer", "Day 3 pain score", "Day 3 visual analogue scale pain score", "0-10", "Copied without change", "Observed source value"],
    ["vas_day_4", "integer or blank", "Day 4 pain score", "Day 4 visual analogue scale pain score", "0-10 or blank", "Copied without change", "Structural blanks remain blank"],
    ["vas_day_5", "integer or blank", "Day 5 pain score", "Day 5 visual analogue scale pain score", "0-10 or blank", "Copied without change", "Structural blanks remain blank"],
    ["vas_day_6", "integer or blank", "Day 6 pain score", "Day 6 visual analogue scale pain score", "0-10 or blank", "Copied without change", "Structural blanks remain blank"],
    ["vas_day_7", "integer or blank", "Day 7 pain score", "Day 7 visual analogue scale pain score", "0-10 or blank", "Copied without change", "No post-cure zeros are derived in this workbook"],
    ["vas_day_28", "integer", "Day 28 pain score", "Day 28 visual analogue scale pain score", "0-10", "Copied without change", "Handled separately from the daily acute-phase sequence"],
    ["pain_duration_days", "integer", "Duration of pain", "Reported pain duration", "2-28 days", "Copied without change", "Value 28 may represent capping or censoring"],
    ["dlqi_admission_total", "integer", "DLQI on admission", "Dermatology Life Quality Index total at admission", "0-30; observed source range 10-20", "Copied without change", "Total score only"],
    ["dlqi_discharge_total", "integer", "DLQI on discharge", "Dermatology Life Quality Index total at discharge", "0-30; observed source range 6-16", "Copied without change", "Total score only"],
    ["nephrotoxicity_flag", "integer", "Nephrotoxicity", "Nephrotoxicity recorded during the study window", "0=No; 1=Yes", "Copied without change", "Onset and resolution dates unavailable"],
    ["hepatotoxicity_flag", "integer", "Hepato-toxicity", "Hepatotoxicity recorded during the study window", "0=No; 1=Yes", "Copied without change", "Onset and resolution dates unavailable"],
    ["nausea_flag", "integer", "Nausea", "Nausea recorded during the study window", "0=No; 1=Yes", "Copied without change", "Onset and resolution dates unavailable"],
    ["vomiting_flag", "integer", "Vomit", "Vomiting recorded during the study window", "0=No; 1=Yes", "Copied without change", "Workbook contains four positive flags"],
    ["upper_gi_bleeding_flag", "integer", "Upper gastrointestinal bleeding", "Upper gastrointestinal bleeding recorded during the study window", "0=No; 1=Yes", "Copied without change", "No positive flags in the workbook"],
    ["fainting_flag", "integer", "Fainting", "Fainting recorded during the study window", "0=No; 1=Yes", "Copied without change", "No positive flags in the workbook"],
    ["needle_break_flag", "integer", "Needle break", "Needle break recorded during the study window", "0=No; 1=Yes", "Copied without change", "No positive flags in the workbook"],
    ["hematoma_flag", "integer", "Hematoma", "Hematoma recorded during the study window", "0=No; 1=Yes", "Copied without change", "No positive flags in the workbook"],
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_full_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", value.strip())
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def parse_discharge_date(value, admission_date: date | None) -> date | None:
    parsed = parse_full_date(value)
    if parsed is not None:
        return parsed
    if admission_date is None:
        return None
    match = re.fullmatch(r"(\d{1,2})\.(\d{1,2})", str(value).strip())
    if not match:
        return None
    month, day = int(match.group(1)), int(match.group(2))
    year = admission_date.year + (1 if month < admission_date.month else 0)
    try:
        return date(year, month, day)
    except ValueError:
        return None


def date_qc_status(row) -> str:
    admission = parse_full_date(row[5])
    discharge = parse_discharge_date(row[6], admission)
    reported_days = row[7]
    if admission is None or discharge is None or not isinstance(reported_days, (int, float)):
        return "UNPARSABLE"
    elapsed = (discharge - admission).days
    if elapsed + 1 == reported_days:
        return "MATCH_INCLUSIVE"
    if elapsed == reported_days:
        return "MATCH_ELAPSED"
    return "DISCREPANT"


def build_output_rows(source_rows):
    output_rows = []
    for row in source_rows:
        group_code = row[8]
        hospitalization_days = row[7]
        if group_code not in (1, 2):
            raise ValueError(f"Unexpected group code for participant {row[0]}")
        if not isinstance(hospitalization_days, int) or hospitalization_days < 1:
            raise ValueError(f"Invalid hospitalization days for participant {row[0]}")
        output_rows.append([
            row[0], row[2], row[3], hospitalization_days, group_code,
            "Control" if group_code == 1 else "Experimental",
            STANDARDIZED_INDEX_DATE,
            STANDARDIZED_INDEX_DATE + timedelta(days=hospitalization_days - 1),
            date_qc_status(row),
            *row[9:29],
        ])
    return output_rows


def style_header(worksheet, fill_color: str) -> None:
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_table(worksheet, reference: str, name: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def create_workbook(source_path: Path, output_path: Path) -> None:
    if not source_path.is_file():
        raise FileNotFoundError(f"Source file not found: {source_path}")
    if output_path.exists():
        raise FileExistsError(f"Output already exists; nothing was overwritten: {output_path}")
    actual_hash = sha256(source_path)
    if actual_hash != EXPECTED_SOURCE_SHA256:
        raise ValueError(f"Source checksum mismatch: {actual_hash}")

    source_workbook = load_workbook(source_path, read_only=True, data_only=True)
    source_sheet = source_workbook[source_workbook.sheetnames[0]]
    source_rows = list(source_sheet.iter_rows(min_row=4, max_row=109, min_col=1, max_col=29, values_only=True))
    if len(source_rows) != 106:
        raise ValueError(f"Expected 106 source records; found {len(source_rows)}")

    output_rows = build_output_rows(source_rows)
    ids = [row[0] for row in output_rows]
    if len(set(ids)) != 106 or min(ids) != 1 or max(ids) != 106:
        raise ValueError("Participant ID reconciliation failed")
    group_counts = Counter(row[4] for row in output_rows)
    if group_counts != Counter({2: 54, 1: 52}):
        raise ValueError(f"Group reconciliation failed: {dict(group_counts)}")
    qc_counts = Counter(row[8] for row in output_rows)
    expected_qc = Counter({"MATCH_INCLUSIVE": 81, "MATCH_ELAPSED": 21, "DISCREPANT": 3, "UNPARSABLE": 1})
    if qc_counts != expected_qc:
        raise ValueError(f"Date reconciliation changed: {dict(qc_counts)}")

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Data"
    dictionary_sheet = workbook.create_sheet("Data_Dictionary")
    log_sheet = workbook.create_sheet("Transformation_Log")

    data_sheet.append(HEADERS)
    for row in output_rows:
        data_sheet.append(row)
    style_header(data_sheet, "1F4E78")
    add_table(data_sheet, "A1:AC107", "HerpesDeidentifiedData")
    data_sheet.freeze_panes = "B2"
    data_sheet.sheet_view.showGridLines = False
    data_sheet.row_dimensions[1].height = 42
    for row in data_sheet.iter_rows(min_row=2, max_row=107, min_col=1, max_col=29):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
    for row_number in range(2, 108):
        data_sheet.cell(row_number, 7).number_format = "yyyy-mm-dd"
        data_sheet.cell(row_number, 8).number_format = "yyyy-mm-dd"
    widths = {"A": 21, "B": 12, "C": 18, "D": 18, "E": 18, "F": 23, "G": 22, "H": 22, "I": 26}
    for column, width in widths.items():
        data_sheet.column_dimensions[column].width = width
    for column_number in range(10, 30):
        data_sheet.column_dimensions[data_sheet.cell(1, column_number).column_letter].width = 19

    dictionary_headers = ["Variable", "Type", "Source", "Definition", "Allowed values or range", "Transformation", "Analysis or OMOP note"]
    dictionary_sheet.append(dictionary_headers)
    for row in DICTIONARY:
        dictionary_sheet.append(row)
    style_header(dictionary_sheet, "4472C4")
    add_table(dictionary_sheet, f"A1:G{len(DICTIONARY) + 1}", "HerpesDataDictionary")
    dictionary_sheet.freeze_panes = "A2"
    dictionary_sheet.sheet_view.showGridLines = False
    for row in dictionary_sheet.iter_rows(min_row=2, max_row=len(DICTIONARY) + 1, min_col=1, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in zip("ABCDEFG", [32, 18, 32, 44, 34, 56, 52]):
        dictionary_sheet.column_dimensions[column].width = width

    transformation_log = [
        ["Source file", source_path.name],
        ["Source SHA-256", actual_hash],
        ["Source classification", "Restricted - pseudonymized"],
        ["Source records", len(source_rows)],
        ["Output records", len(output_rows)],
        ["Source columns", 29],
        ["Output columns", len(HEADERS)],
        ["Excluded source fields", "Name; Admission number; Admission time; Discharge time"],
        ["Name-column clarification", "The single nonblank cell is an explanatory note; no participant name is present or retained."],
        ["Construction method", "New workbook created from an approved-field allowlist; the source workbook was not copied or edited."],
        ["Standardized index date", "2000-01-01 for all participants"],
        ["Standardized discharge rule", "standardized_index_date + hospitalization_days_source - 1 day"],
        ["Date interval QC - MATCH_INCLUSIVE", qc_counts["MATCH_INCLUSIVE"]],
        ["Date interval QC - MATCH_ELAPSED", qc_counts["MATCH_ELAPSED"]],
        ["Date interval QC - DISCREPANT", qc_counts["DISCREPANT"]],
        ["Date interval QC - UNPARSABLE", qc_counts["UNPARSABLE"]],
        ["Group reconciliation", f"Control={group_counts[1]}; Experimental={group_counts[2]}"],
        ["VAS missingness rule", "Source blanks preserved. No structural zeros are inserted in this working copy."],
        ["Intended use", "Deidentified staging, source replication, OMOP ETL development, and analysis."],
        ["Version", VERSION],
    ]
    log_sheet.append(["Item", "Value"])
    for row in transformation_log:
        log_sheet.append(row)
    style_header(log_sheet, "5B9BD5")
    add_table(log_sheet, f"A1:B{len(transformation_log) + 1}", "HerpesTransformationLog")
    log_sheet.freeze_panes = "A2"
    log_sheet.sheet_view.showGridLines = False
    for row in log_sheet.iter_rows(min_row=2, max_row=len(transformation_log) + 1, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    log_sheet.column_dimensions["A"].width = 42
    log_sheet.column_dimensions["B"].width = 100

    workbook.properties.creator = "QiBo Institute"
    workbook.properties.title = "Herpes Pain Deidentified Working Data"
    workbook.properties.subject = "Study replication staging data"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    print("DEIDENTIFICATION_STATUS=PASS")
    print(f"SOURCE_SHA256={actual_hash}")
    print(f"OUTPUT_SHA256={sha256(output_path)}")
    print(f"OUTPUT_RECORDS={len(output_rows)}")
    print(f"GROUP_COUNTS={dict(sorted(group_counts.items()))}")
    print(f"DATE_QC_COUNTS={dict(sorted(qc_counts.items()))}")
    print(f"OUTPUT={output_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path, help="Restricted source XLSX")
    parser.add_argument("--output", required=True, type=Path, help="New deidentified XLSX")
    args = parser.parse_args()
    try:
        create_workbook(args.source.resolve(), args.output.resolve())
    except (FileNotFoundError, FileExistsError, ValueError, OSError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
